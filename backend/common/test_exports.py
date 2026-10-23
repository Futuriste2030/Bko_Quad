"""Tests des exports CSV/XLSX : contenu, formats, rôles et site."""

import csv
from datetime import timedelta
from io import BytesIO, StringIO

from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import User
from bookings.models import Booking
from catalogue.models import Activity, Site, Slot
from clients.models import Customer
from finance.models import BankTransaction, CashbookEntry
from payments.models import Invoice, Payment
from payroll.models import Employee, PayrollSheet, Payslip
from telephony.models import CallLog


class ExportApiTest(TestCase):
    def setUp(self):
        today = timezone.localdate()
        start = timezone.now()
        self.site_a = Site.objects.create(name="Orpheus", city="Bamako")
        self.site_b = Site.objects.create(name="Konibabougou", city="Bamako")

        self.act_a = Activity.objects.create(
            site=self.site_a, name="Quad 1h", duration_minutes=60,
            base_price=25000, deposit=10000, max_group=10)
        self.act_b = Activity.objects.create(
            site=self.site_b, name="Quad 2h", duration_minutes=120,
            base_price=45000, deposit=15000, max_group=10)
        self.slot_a = Slot.objects.create(
            activity=self.act_a, start=start, end=start + timedelta(hours=1),
            capacity=10)
        self.slot_b = Slot.objects.create(
            activity=self.act_b, start=start, end=start + timedelta(hours=2),
            capacity=10)

        self.agent_a = User.objects.create_user(
            username="agentA", password="x", phone="+22370000001",
            role="agent", site=self.site_a)
        self.agent_b = User.objects.create_user(
            username="agentB", password="x", phone="+22370000002",
            role="agent", site=self.site_b)
        self.resp_a = User.objects.create_user(
            username="respA", password="x", phone="+22370000003",
            role="responsable", site=self.site_a)
        self.comptable = User.objects.create_user(
            username="compta", password="x", phone="+22370000004",
            role="comptable")
        self.admin = User.objects.create_user(
            username="boss", password="x", phone="+22370000005",
            role="admin")
        self.guide = User.objects.create_user(
            username="guideA", password="x", phone="+22370000006",
            role="guide", site=self.site_a)

        self.customer_a = Customer.objects.create(
            name="Client A", phone="+22376000001")
        self.customer_b = Customer.objects.create(
            name="Client B", phone="+22376000002")
        self.booking_a = self._booking(
            "BQ-A", self.customer_a, self.slot_a, self.agent_a, 25000)
        self.booking_b = self._booking(
            "BQ-B", self.customer_b, self.slot_b, self.agent_b, 45000)

        self.payment_a = Payment.objects.create(
            booking=self.booking_a, nature="deposit", mode="cash",
            amount=10000, reference="PAY-A", status="paid",
            received_by=self.agent_a, paid_at=timezone.now())
        self.payment_b = Payment.objects.create(
            booking=self.booking_b, nature="deposit", mode="cash",
            amount=15000, reference="PAY-B", status="paid",
            received_by=self.agent_b, paid_at=timezone.now())
        self.invoice_a = Invoice.objects.create(booking=self.booking_a, number=1)
        self.invoice_b = Invoice.objects.create(booking=self.booking_b, number=2)

        CashbookEntry.objects.create(
            date=today, site=self.site_a, category="carburant",
            label="Carburant A", voucher="CASH-A", direction="out",
            amount=5000, created_by=self.resp_a)
        CashbookEntry.objects.create(
            date=today, site=self.site_b, category="carburant",
            label="Carburant B", voucher="CASH-B", direction="out",
            amount=7000, created_by=self.resp_a)

        BankTransaction.objects.create(
            type="deposit", date=today, reason="Versement A", amount=50000,
            bank="BDM", person_first_name="Awa", person_last_name="Diarra",
            receipt_number="BORD-A", site=self.site_a,
            created_by=self.comptable)
        BankTransaction.objects.create(
            type="withdrawal", date=today, reason="Retrait B", amount=20000,
            bank="BDM", person_first_name="Moussa", person_last_name="Keita",
            check_number="CHQ-B", site=self.site_b, created_by=self.comptable)

        CallLog.objects.create(
            direction="inbound", phone_number="+22376000001",
            started_at=timezone.now(), duration_seconds=60,
            external_id="CALL-A")
        CallLog.objects.create(
            direction="missed", phone_number="+22376000002",
            started_at=timezone.now(), duration_seconds=0,
            external_id="CALL-B")

        sheet = PayrollSheet.objects.create(month=today.replace(day=1))
        employee_a = Employee.objects.create(
            name="Agent A", role="agent", site=self.site_a)
        employee_b = Employee.objects.create(
            name="Agent B", role="agent", site=self.site_b)
        Payslip.objects.create(
            sheet=sheet, employee=employee_a, site_name=self.site_a.name,
            base_amount=100000, gross=100000, deductions=0, net=100000)
        Payslip.objects.create(
            sheet=sheet, employee=employee_b, site_name=self.site_b.name,
            base_amount=120000, gross=120000, deductions=0, net=120000)

        self.start = today.isoformat()
        self.end = today.isoformat()

    def _booking(self, code, customer, slot, user, total):
        return Booking.objects.create(
            code=code, customer=customer, slot=slot, persons=1, quads=1,
            unit_price=total, total=total, deposit_due=total // 2,
            qr_token=f"qr-{code}", created_by=user,
            expires_at=timezone.now() + timedelta(minutes=20))

    def _client(self, user):
        client = APIClient()
        client.force_authenticate(user=user)
        return client

    def _get(self, user, dataset, file_type="csv", **params):
        query = {
            "dataset": dataset,
            "file": file_type,
            "start": self.start,
            "end": self.end,
            **{key: str(value) for key, value in params.items()},
        }
        return self._client(user).get("/api/v1/exports/", query)

    def _csv_rows(self, response):
        text = response.content.decode("utf-8-sig")
        return list(csv.reader(StringIO(text), delimiter=";"))

    def test_csv_admin_reservations(self):
        response = self._get(self.admin, "reservations")
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("attachment; filename=", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        rows = self._csv_rows(response)
        self.assertEqual(rows[0][0], "Date créneau")
        self.assertIn("BQ-A", [row[1] for row in rows[1:]])
        self.assertIn("BQ-B", [row[1] for row in rows[1:]])

    def test_agent_site_scope_and_finance_refusal(self):
        response = self._get(self.agent_a, "reservations")
        self.assertEqual(response.status_code, 200)
        rows = self._csv_rows(response)
        self.assertEqual([row[1] for row in rows[1:]], ["BQ-A"])

        self.assertEqual(
            self._get(self.agent_a, "reservations", site=self.site_b.id).status_code,
            403)
        self.assertEqual(
            self._get(self.agent_a, "paiements").status_code, 403)
        self.assertEqual(
            self._get(self.agent_a, "caisse").status_code, 403)

    def test_comptable_payments_and_site_filter(self):
        response = self._get(self.comptable, "paiements")
        self.assertEqual(response.status_code, 200)
        rows = self._csv_rows(response)
        self.assertEqual({row[1] for row in rows[1:]}, {"PAY-A", "PAY-B"})

        response = self._get(
            self.comptable, "paiements", site=self.site_a.id)
        rows = self._csv_rows(response)
        self.assertEqual([row[1] for row in rows[1:]], ["PAY-A"])

    def test_banque_filtre_par_site(self):
        response = self._get(self.admin, "banque", site=self.site_a.id)
        self.assertEqual(response.status_code, 200)
        rows = self._csv_rows(response)
        self.assertEqual(rows[0][2], "Site")
        self.assertEqual([row[2] for row in rows[1:]], ["Orpheus"])

    def test_activity_agent_and_calls(self):
        response = self._get(
            self.admin, "reservations", activity=self.act_b.id)
        self.assertEqual([row[1] for row in self._csv_rows(response)[1:]], ["BQ-B"])

        response = self._get(
            self.admin, "reservations", agent="agentA")
        self.assertEqual([row[1] for row in self._csv_rows(response)[1:]], ["BQ-A"])

        response = self._get(self.admin, "appels")
        self.assertEqual(response.status_code, 200)
        rows = self._csv_rows(response)
        self.assertEqual({row[2] for row in rows[1:]},
                         {"+22376000001", "+22376000002"})

    def test_xlsx_payroll_and_site_filter(self):
        response = self._get(self.comptable, "paie", file_type="xlsx")
        self.assertEqual(response.status_code, 200, response.content[:300])
        self.assertIn("spreadsheetml.sheet", response["Content-Type"])
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        self.assertEqual(sheet["A1"].value, "Numéro")
        self.assertEqual(sheet.max_row, 3)

        response = self._get(
            self.comptable, "paie", file_type="xlsx", site=self.site_b.id)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet["C2"].value, "Agent B")

    def test_invalid_parameters_and_permissions(self):
        self.assertEqual(
            self._client(self.admin).get("/api/v1/exports/").status_code, 400)
        self.assertEqual(
            self._get(self.admin, "inconnu").status_code, 400)
        self.assertEqual(
            self._get(self.admin, "reservations", file_type="pdf").status_code,
            400)
        self.assertEqual(
            self._get(self.admin, "reservations", start="2026-10-20",
                     end="2026-10-01").status_code, 400)
        self.assertEqual(
            self._get(self.guide, "reservations").status_code, 403)
        self.assertEqual(
            self._get(self.admin, "banque", activity=self.act_a.id).status_code,
            400)
        self.assertEqual(
            self._get(self.admin, "appels", agent="agentA").status_code, 400)
