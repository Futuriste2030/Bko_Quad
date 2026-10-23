"""Exports CSV/XLSX des données d'exploitation et de gestion.

Chaque jeu de données est filtré par période et, quand il le permet, par
site, activité et agent. Les permissions et le site imposé sont vérifiés
par la vue avant l'appel des constructeurs.
"""

import csv
import re
from dataclasses import dataclass
from datetime import date
from io import BytesIO, StringIO
from typing import Callable, Iterable

from django.db.models import Q, QuerySet, Sum
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bookings.models import Booking
from finance.models import BankTransaction, CashbookEntry
from payments.models import Invoice, Payment
from payroll.models import Payslip
from telephony.models import CallLog


class ExportError(ValueError):
    pass


@dataclass(frozen=True)
class ExportFilters:
    start: date
    end: date
    site: int | None = None
    activity: int | None = None
    agent: str = ""


@dataclass(frozen=True)
class ExportDataset:
    key: str
    label: str
    roles: tuple[str, ...]
    columns: tuple[str, ...]
    build: Callable[[date, date], QuerySet]
    row: Callable[[object], list]
    site_lookup: str = ""
    activity_lookup: str = ""
    agent_lookups: tuple[str, ...] = ()


def _date(value):
    return value.strftime("%d/%m/%Y") if value else ""


def _datetime(value):
    if not value:
        return ""
    if timezone.is_aware(value):
        value = timezone.localtime(value)
    return value.strftime("%d/%m/%Y %H:%M")


def _name(user):
    if not user:
        return ""
    return user.get_full_name() or user.username


def _payments(start, end):
    return (
        Payment.objects.filter(
            Q(paid_at__date__range=(start, end))
            | Q(paid_at__isnull=True, created_at__date__range=(start, end))
        )
        .select_related(
            "booking__customer",
            "booking__slot__activity__site",
            "received_by",
        )
        .order_by("created_at", "id")
    )


def _payslips(start, end):
    first = start.replace(day=1)
    last = end.replace(day=1)
    return (
        Payslip.objects.filter(sheet__month__gte=first, sheet__month__lte=last)
        .annotate(paid=Sum("payments__amount"))
        .select_related("sheet", "employee", "employee__site", "user")
        .order_by("sheet__month", "employee__name", "id")
    )


def _reservation_row(obj: Booking):
    slot = obj.slot
    activity = slot.activity
    return [
        _date(slot.start),
        obj.code,
        obj.customer.name,
        obj.customer.phone,
        activity.site.name if activity.site_id else "",
        activity.name,
        _datetime(slot.start),
        _datetime(slot.end),
        obj.persons,
        obj.quads,
        obj.unit_price,
        obj.units,
        obj.total,
        obj.deposit_due,
        obj.get_status_display(),
        obj.get_source_display(),
        _name(obj.created_by),
        _datetime(obj.created_at),
        obj.cancel_reason,
    ]


def _payment_row(obj: Payment):
    booking = obj.booking
    activity = booking.slot.activity
    return [
        _datetime(obj.paid_at),
        obj.reference,
        booking.code,
        booking.customer.name,
        activity.site.name if activity.site_id else "",
        activity.name,
        obj.get_nature_display(),
        obj.get_mode_display(),
        obj.amount,
        obj.get_status_display(),
        _name(obj.received_by),
        _datetime(obj.created_at),
    ]


def _cash_row(obj: CashbookEntry):
    return [
        _date(obj.date),
        obj.voucher,
        obj.site.name if obj.site_id else "",
        obj.get_direction_display(),
        obj.get_category_display(),
        obj.label,
        obj.amount,
        obj.booking.code if obj.booking_id else "",
        obj.employee_name,
        obj.receipt,
        "Oui" if obj.locked else "Non",
        _name(obj.created_by),
    ]


def _bank_row(obj: BankTransaction):
    credit = obj.amount if obj.type in (
        BankTransaction.DEPOSIT, BankTransaction.TRANSFER) else 0
    debit = obj.amount if obj.type == BankTransaction.WITHDRAWAL else 0
    return [
        _date(obj.date),
        obj.get_type_display(),
        obj.site.name if obj.site_id else "",
        obj.bank,
        debit,
        credit,
        obj.reason,
        obj.receipt_number,
        obj.check_number,
        obj.counterparty_bank,
        f"{obj.person_first_name} {obj.person_last_name}".strip(),
        _name(obj.created_by),
        _datetime(obj.created_at),
    ]


def _invoice_row(obj: Invoice):
    booking = obj.booking
    activity = booking.slot.activity
    return [
        f"{obj.number:06d}",
        _datetime(obj.issued_at),
        booking.code,
        booking.customer.name,
        activity.site.name if activity.site_id else "",
        activity.name,
        booking.total,
        booking.get_status_display(),
    ]


def _payslip_row(obj: Payslip):
    paid = getattr(obj, "paid", 0) or 0
    return [
        obj.number,
        obj.sheet.month.strftime("%m/%Y"),
        obj.employee.name,
        obj.position,
        obj.site_name,
        obj.base_amount,
        obj.gross,
        obj.deductions,
        obj.net,
        paid,
        max(0, obj.net - paid),
        obj.sheet.get_status_display(),
    ]


def _call_row(obj: CallLog):
    anomaly = getattr(obj, "anomaly", None)
    return [
        _datetime(obj.started_at),
        obj.get_direction_display(),
        obj.phone_number,
        obj.duration_seconds,
        obj.source,
        "Oui" if anomaly else "Non",
        anomaly.get_reason_display() if anomaly and anomaly.reason else "",
        anomaly.comment if anomaly else "",
        _datetime(anomaly.resolved_at) if anomaly else "",
    ]


DATASETS = {
    "reservations": ExportDataset(
        key="reservations",
        label="Réservations",
        roles=("agent", "responsable", "admin", "comptable"),
        columns=(
            "Date créneau", "Code", "Client", "Téléphone", "Site", "Activité",
            "Début", "Fin", "Personnes", "Quads", "Prix unitaire (FCFA)",
            "Unités", "Total (FCFA)", "Acompte (FCFA)", "Statut", "Source",
            "Agent", "Créée le", "Motif d'annulation",
        ),
        build=lambda start, end: (
            Booking.objects.filter(
                slot__start__date__range=(start, end))
            .select_related(
                "customer", "slot__activity__site", "created_by")
            .order_by("slot__start", "id")
        ),
        row=_reservation_row,
        site_lookup="slot__activity__site",
        activity_lookup="slot__activity",
        agent_lookups=(
            "created_by__username", "created_by__first_name",
            "created_by__last_name"),
    ),
    "paiements": ExportDataset(
        key="paiements",
        label="Paiements",
        roles=("responsable", "admin", "comptable"),
        columns=(
            "Date paiement", "Référence", "Réservation", "Client", "Site",
            "Activité", "Nature", "Mode", "Montant (FCFA)", "Statut",
            "Encaissé par", "Créé le",
        ),
        build=_payments,
        row=_payment_row,
        site_lookup="booking__slot__activity__site",
        activity_lookup="booking__slot__activity",
        agent_lookups=("received_by__username", "received_by__first_name",
                       "received_by__last_name"),
    ),
    "caisse": ExportDataset(
        key="caisse",
        label="Caisse",
        roles=("responsable", "admin", "comptable"),
        columns=(
            "Date", "Bordereau", "Site", "Sens", "Catégorie", "Libellé",
            "Montant (FCFA)", "Réservation", "Employé", "Justificatif",
            "Verrouillée", "Saisie par",
        ),
        build=lambda start, end: (
            CashbookEntry.objects.filter(date__range=(start, end))
            .select_related("site", "booking", "created_by")
            .order_by("date", "id")
        ),
        row=_cash_row,
        site_lookup="site",
        activity_lookup="booking__slot__activity",
        agent_lookups=("created_by__username", "created_by__first_name",
                       "created_by__last_name"),
    ),
    "banque": ExportDataset(
        key="banque",
        label="Banque",
        roles=("responsable", "admin", "comptable"),
        columns=(
            "Date", "Type", "Site", "Banque", "Débit (FCFA)", "Crédit (FCFA)",
            "Motif", "Bordereau", "Chèque / référence",
            "Banque contrepartie", "Personne", "Saisie par", "Créée le",
        ),
        build=lambda start, end: (
            BankTransaction.objects.filter(date__range=(start, end))
            .select_related("site", "created_by")
            .order_by("date", "id")
        ),
        row=_bank_row,
        site_lookup="site",
        agent_lookups=("created_by__username", "created_by__first_name",
                       "created_by__last_name"),
    ),
    "factures": ExportDataset(
        key="factures",
        label="Factures",
        roles=("agent", "responsable", "admin", "comptable"),
        columns=(
            "Numéro", "Date", "Réservation", "Client", "Site", "Activité",
            "Total réservation (FCFA)", "Statut réservation",
        ),
        build=lambda start, end: (
            Invoice.objects.filter(issued_at__date__range=(start, end))
            .select_related(
                "booking__customer", "booking__slot__activity__site",
                "booking__created_by")
            .order_by("issued_at", "id")
        ),
        row=_invoice_row,
        site_lookup="booking__slot__activity__site",
        activity_lookup="booking__slot__activity",
        agent_lookups=("booking__created_by__username",
                       "booking__created_by__first_name",
                       "booking__created_by__last_name"),
    ),
    "paie": ExportDataset(
        key="paie",
        label="Paie",
        roles=("admin", "comptable"),
        columns=(
            "Numéro", "Mois", "Employé", "Poste", "Site", "Base (FCFA)",
            "Brut (FCFA)", "Retenues (FCFA)", "Net (FCFA)",
            "Versé (FCFA)", "Reste (FCFA)", "Fiche",
        ),
        build=_payslips,
        row=_payslip_row,
        site_lookup="employee__site",
        agent_lookups=("employee__name", "employee__user__username",
                       "employee__user__first_name",
                       "employee__user__last_name"),
    ),
    "appels": ExportDataset(
        key="appels",
        label="Appels",
        roles=("agent", "responsable", "admin"),
        columns=(
            "Date", "Sens", "Numéro", "Durée (s)", "Source", "Anomalie",
            "Motif", "Commentaire", "Résolue le",
        ),
        build=lambda start, end: (
            CallLog.objects.filter(started_at__date__range=(start, end))
            .select_related("anomaly")
            .order_by("started_at", "id")
        ),
        row=_call_row,
    ),
}


def get_dataset(key: str) -> ExportDataset:
    dataset = DATASETS.get((key or "").strip().lower())
    if dataset is None:
        raise ExportError("Jeu de données inconnu.")
    return dataset


def can_export(user, dataset: ExportDataset) -> bool:
    if getattr(user, "is_superuser", False):
        return True
    role = getattr(user, "role", None)
    return role == "superadmin" or role in dataset.roles


def apply_filters(dataset: ExportDataset, filters: ExportFilters):
    qs = dataset.build(filters.start, filters.end)
    if dataset.site_lookup and filters.site:
        qs = qs.filter(**{dataset.site_lookup: filters.site})
    if dataset.activity_lookup and filters.activity:
        qs = qs.filter(**{dataset.activity_lookup: filters.activity})
    if dataset.agent_lookups and filters.agent:
        query = Q()
        for field in dataset.agent_lookups:
            query |= Q(**{f"{field}__icontains": filters.agent})
        qs = qs.filter(query)
    return qs


def _csv_bytes(dataset: ExportDataset, rows: Iterable[list]) -> bytes:
    buffer = StringIO(newline="")
    buffer.write("\ufeff")
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(dataset.columns)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def _xlsx_bytes(dataset: ExportDataset, rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    title = re.sub(r"[:\\/?*\[\]]", "-", dataset.label)[:31] or "Export"
    sheet.title = title
    sheet.sheet_view.showGridLines = False
    sheet.append(list(dataset.columns))

    orange = PatternFill("solid", fgColor="E8590C")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="0A0A0A")
        cell.fill = orange
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for row in rows:
        sheet.append(row)

    for index, header in enumerate(dataset.columns, 1):
        letter = get_column_letter(index)
        values = [str(row[index - 1]) for row in rows]
        width = max([len(str(header)) + 2] + [len(value) + 2 for value in values])
        sheet.column_dimensions[letter].width = min(48, max(12, width))
        if "(FCFA)" in header:
            for cell in sheet[letter][1:]:
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")

    last = f"{get_column_letter(len(dataset.columns))}{len(rows) + 1}"
    sheet.auto_filter.ref = f"A1:{last}"
    sheet.freeze_panes = "A2"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_export(dataset: ExportDataset, filters: ExportFilters, file_type: str):
    rows = [dataset.row(obj) for obj in apply_filters(dataset, filters)]
    if file_type == "csv":
        return _csv_bytes(dataset, rows), "text/csv; charset=utf-8"
    if file_type == "xlsx":
        return (
            _xlsx_bytes(dataset, rows),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    raise ExportError("Format inconnu (csv ou xlsx).")
